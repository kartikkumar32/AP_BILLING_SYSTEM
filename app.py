import os, io, sqlite3, secrets
from datetime import datetime, date
from functools import wraps
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file, abort
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.middleware.proxy_fix import ProxyFix
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook

APP_NAME = "Arvind Pandey Food and Catering Services"
app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1)
DB_PATH = os.environ.get("DATABASE_PATH", os.path.join(os.path.dirname(__file__), "billing.db"))
GSTIN = os.environ.get("BUSINESS_GSTIN", "")
BUSINESS_ADDRESS = os.environ.get("BUSINESS_ADDRESS", "")
BUSINESS_PHONE = os.environ.get("BUSINESS_PHONE", "")

def db():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys=ON")
    return con

def init_db():
    with db() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, username TEXT UNIQUE NOT NULL, password_hash TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('owner','staff')), active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS customers(id INTEGER PRIMARY KEY, name TEXT NOT NULL, phone TEXT, email TEXT, address TEXT, gstin TEXT, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS invoices(id INTEGER PRIMARY KEY, invoice_no TEXT UNIQUE NOT NULL, customer_id INTEGER NOT NULL REFERENCES customers(id), event_name TEXT, event_date TEXT, billing_mode TEXT NOT NULL, subtotal_paise INTEGER NOT NULL, discount_paise INTEGER NOT NULL DEFAULT 0, tax_rate REAL NOT NULL DEFAULT 0, tax_paise INTEGER NOT NULL DEFAULT 0, total_paise INTEGER NOT NULL, notes TEXT, created_by INTEGER REFERENCES users(id), created_at TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'issued');
        CREATE TABLE IF NOT EXISTS invoice_items(id INTEGER PRIMARY KEY, invoice_id INTEGER NOT NULL REFERENCES invoices(id) ON DELETE CASCADE, description TEXT NOT NULL, qty REAL NOT NULL, unit TEXT NOT NULL, rate_paise INTEGER NOT NULL, amount_paise INTEGER NOT NULL);
        CREATE TABLE IF NOT EXISTS payments(id INTEGER PRIMARY KEY, invoice_id INTEGER NOT NULL REFERENCES invoices(id), amount_paise INTEGER NOT NULL CHECK(amount_paise>0), method TEXT NOT NULL, reference TEXT, paid_on TEXT NOT NULL, notes TEXT, created_by INTEGER REFERENCES users(id), created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS expenses(id INTEGER PRIMARY KEY, expense_date TEXT NOT NULL, category TEXT NOT NULL, description TEXT, amount_paise INTEGER NOT NULL CHECK(amount_paise>=0), payment_method TEXT, vendor TEXT, created_by INTEGER REFERENCES users(id), created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS inventory(id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE, unit TEXT NOT NULL, quantity REAL NOT NULL DEFAULT 0, reorder_level REAL NOT NULL DEFAULT 0, unit_cost_paise INTEGER NOT NULL DEFAULT 0, updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS stock_moves(id INTEGER PRIMARY KEY, item_id INTEGER NOT NULL REFERENCES inventory(id), delta REAL NOT NULL, reason TEXT NOT NULL, note TEXT, created_by INTEGER REFERENCES users(id), created_at TEXT NOT NULL);
        CREATE INDEX IF NOT EXISTS idx_invoice_customer ON invoices(customer_id);
        CREATE INDEX IF NOT EXISTS idx_invoice_date ON invoices(created_at);
        CREATE INDEX IF NOT EXISTS idx_payments_invoice ON payments(invoice_id);
        """)
        if not c.execute("SELECT id FROM users LIMIT 1").fetchone():
            username = os.environ.get("OWNER_USERNAME", "owner")
            password = os.environ.get("OWNER_PASSWORD")
            if not password:
                password = secrets.token_urlsafe(12)
                print("\\nINITIAL OWNER LOGIN (save this now):", username, password, "\\nSet OWNER_PASSWORD and SECRET_KEY in production environment.\\n")
            c.execute("INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)",
                      (username, generate_password_hash(password), "owner", datetime.now().isoformat(timespec="seconds")))

init_db()

def money(paise):
    # Use Rs. because ReportLab Helvetica may not contain the Unicode rupee glyph.
    return f"Rs. {(int(paise or 0)/100):,.2f}"
app.jinja_env.filters["money"] = money

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login", next=request.path))
        with db() as c:
            u = c.execute("SELECT * FROM users WHERE id=? AND active=1", (session["user_id"],)).fetchone()
        if not u:
            session.clear()
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapped

def owner_required(fn):
    @wraps(fn)
    @login_required
    def wrapped(*args, **kwargs):
        if session.get("role") != "owner":
            abort(403)
        return fn(*args, **kwargs)
    return wrapped

def to_paise(value):
    try:
        d = Decimal(str(value or "0"))
        if d < 0: raise ValueError("Amounts cannot be negative.")
        return int((d * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError):
        raise ValueError("Enter a valid non-negative amount.")

def safe_float(value, minimum=0):
    try:
        x = float(value)
        if x < minimum: raise ValueError
        return x
    except (TypeError, ValueError):
        raise ValueError("Enter a valid quantity.")

@app.after_request
def security_headers(resp):
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    resp.headers["Cache-Control"] = "no-store"
    return resp

@app.context_processor
def inject_globals():
    return {"business_name": APP_NAME, "business_address": BUSINESS_ADDRESS, "business_phone": BUSINESS_PHONE,
            "business_gstin": GSTIN, "current_user": session.get("username"), "current_role": session.get("role")}

@app.route("/")
@login_required
def dashboard():
    with db() as c:
        invoices = c.execute("SELECT COUNT(*) n, COALESCE(SUM(total_paise),0) total FROM invoices WHERE status='issued'").fetchone()
        received = c.execute("SELECT COALESCE(SUM(amount_paise),0) total FROM payments").fetchone()["total"]
        expense = c.execute("SELECT COALESCE(SUM(amount_paise),0) total FROM expenses").fetchone()["total"]
        pending = c.execute("SELECT COALESCE(SUM(i.total_paise),0)-COALESCE((SELECT SUM(p.amount_paise) FROM payments p WHERE p.invoice_id=i.id),0) bal FROM invoices i WHERE i.status='issued'").fetchone()["bal"]
        recent = c.execute("""SELECT i.*, c.name customer FROM invoices i JOIN customers c ON c.id=i.customer_id ORDER BY i.id DESC LIMIT 8""").fetchall()
        lowstock = c.execute("SELECT * FROM inventory WHERE quantity<=reorder_level ORDER BY name LIMIT 6").fetchall()
    return render_template("dashboard.html", invoice_count=invoices["n"], sales=invoices["total"], received=received, pending=max(0,pending), expense=expense, recent=recent, lowstock=lowstock)

@app.route("/login", methods=["GET","POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username","").strip()
        password = request.form.get("password","")
        with db() as c:
            user = c.execute("SELECT * FROM users WHERE username=? AND active=1", (username,)).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            session.clear()
            session["user_id"], session["username"], session["role"] = user["id"], user["username"], user["role"]
            return redirect(request.args.get("next") or url_for("dashboard"))
        flash("Invalid username or password.", "danger")
    return render_template("login.html")

@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/customers", methods=["GET","POST"])
@login_required
def customers():
    if request.method == "POST":
        name = request.form.get("name","").strip()
        if not name:
            flash("Customer name is required.", "danger")
        else:
            with db() as c:
                c.execute("INSERT INTO customers(name,phone,email,address,gstin,created_at) VALUES(?,?,?,?,?,?)",
                          (name, request.form.get("phone","").strip(), request.form.get("email","").strip(),
                           request.form.get("address","").strip(), request.form.get("gstin","").strip(), datetime.now().isoformat(timespec="seconds")))
            flash("Customer saved.", "success")
        return redirect(url_for("customers"))
    with db() as c:
        rows = c.execute("SELECT * FROM customers ORDER BY name").fetchall()
    return render_template("customers.html", customers=rows)

@app.route("/customers/<int:customer_id>/ledger")
@login_required
def customer_ledger(customer_id):
    with db() as c:
        customer = c.execute("SELECT * FROM customers WHERE id=?", (customer_id,)).fetchone()
        if not customer: abort(404)
        invoices = c.execute("""SELECT i.*, COALESCE((SELECT SUM(p.amount_paise) FROM payments p WHERE p.invoice_id=i.id),0) paid
            FROM invoices i WHERE customer_id=? ORDER BY i.id DESC""", (customer_id,)).fetchall()
    return render_template("ledger.html", customer=customer, invoices=invoices)

@app.route("/invoices")
@login_required
def invoices():
    with db() as c:
        rows = c.execute("""SELECT i.*, c.name customer,
          COALESCE((SELECT SUM(p.amount_paise) FROM payments p WHERE p.invoice_id=i.id),0) paid
          FROM invoices i JOIN customers c ON c.id=i.customer_id ORDER BY i.id DESC""").fetchall()
    return render_template("invoices.html", invoices=rows)

@app.route("/invoices/new", methods=["GET","POST"])
@login_required
def invoice_new():
    with db() as c:
        customer_rows = c.execute("SELECT * FROM customers ORDER BY name").fetchall()
    if request.method == "POST":
        try:
            customer_id = int(request.form["customer_id"])
            mode = request.form.get("billing_mode","items")
            if mode not in ("items","plates"): raise ValueError("Choose a valid billing mode.")
            descs = request.form.getlist("description[]")
            qtys = request.form.getlist("qty[]")
            units = request.form.getlist("unit[]")
            rates = request.form.getlist("rate[]")
            items = []
            subtotal = 0
            for d,q,u,r in zip(descs,qtys,units,rates):
                if not d.strip(): continue
                qty = safe_float(q, 0.000001)
                rate = to_paise(r)
                amount = int(Decimal(str(qty * rate)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
                items.append((d.strip(),qty,u.strip() or ("plate" if mode=="plates" else "item"),rate,amount))
                subtotal += amount
            if not items: raise ValueError("Add at least one line item.")
            discount = to_paise(request.form.get("discount","0"))
            if discount > subtotal: raise ValueError("Discount cannot exceed subtotal.")
            tax_rate = float(request.form.get("tax_rate","0") or 0)
            if tax_rate < 0 or tax_rate > 100: raise ValueError("Tax rate must be between 0 and 100.")
            taxable = subtotal - discount
            tax = int((Decimal(taxable) * Decimal(str(tax_rate)) / Decimal(100)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
            total = taxable + tax
            with db() as c:
                customer = c.execute("SELECT id FROM customers WHERE id=?", (customer_id,)).fetchone()
                if not customer: raise ValueError("Select a valid customer.")
                stamp = datetime.now()
                invno = f"APCS-{stamp:%Y%m%d}-{secrets.token_hex(3).upper()}"
                cur = c.execute("""INSERT INTO invoices(invoice_no,customer_id,event_name,event_date,billing_mode,subtotal_paise,discount_paise,tax_rate,tax_paise,total_paise,notes,created_by,created_at)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (invno, customer_id, request.form.get("event_name","").strip(), request.form.get("event_date") or None, mode,
                     subtotal, discount, tax_rate, tax, total, request.form.get("notes","").strip(), session["user_id"], stamp.isoformat(timespec="seconds")))
                iid = cur.lastrowid
                c.executemany("INSERT INTO invoice_items(invoice_id,description,qty,unit,rate_paise,amount_paise) VALUES(?,?,?,?,?,?)",
                              [(iid,*it) for it in items])
            flash(f"Invoice {invno} created.", "success")
            return redirect(url_for("invoice_detail", invoice_id=iid))
        except (ValueError, KeyError, sqlite3.Error) as e:
            flash(str(e) if isinstance(e, ValueError) else "Could not save invoice. Check the form and try again.", "danger")
    return render_template("invoice_new.html", customers=customer_rows)

@app.route("/invoices/<int:invoice_id>")
@login_required
def invoice_detail(invoice_id):
    with db() as c:
        inv = c.execute("SELECT i.*, c.name customer,c.phone,c.email,c.address,c.gstin customer_gstin FROM invoices i JOIN customers c ON c.id=i.customer_id WHERE i.id=?", (invoice_id,)).fetchone()
        if not inv: abort(404)
        items = c.execute("SELECT * FROM invoice_items WHERE invoice_id=?", (invoice_id,)).fetchall()
        payments = c.execute("SELECT * FROM payments WHERE invoice_id=? ORDER BY paid_on DESC,id DESC", (invoice_id,)).fetchall()
    paid = sum(p["amount_paise"] for p in payments)
    return render_template("invoice_detail.html", inv=inv, items=items, payments=payments, paid=paid, balance=max(0,inv["total_paise"]-paid))

@app.route("/invoices/<int:invoice_id>/payments", methods=["POST"])
@login_required
def payment_add(invoice_id):
    try:
        amount = to_paise(request.form.get("amount"))
        if amount <= 0: raise ValueError("Payment must be greater than zero.")
        method = request.form.get("method","UPI")
        if method not in ("Cash","UPI","Bank transfer","Card","Cheque","Other"): raise ValueError("Invalid payment method.")
        with db() as c:
            inv = c.execute("SELECT total_paise FROM invoices WHERE id=?", (invoice_id,)).fetchone()
            if not inv: abort(404)
            paid = c.execute("SELECT COALESCE(SUM(amount_paise),0) n FROM payments WHERE invoice_id=?", (invoice_id,)).fetchone()["n"]
            if amount > inv["total_paise"] - paid: raise ValueError("Payment cannot exceed the outstanding balance.")
            c.execute("INSERT INTO payments(invoice_id,amount_paise,method,reference,paid_on,notes,created_by,created_at) VALUES(?,?,?,?,?,?,?,?)",
                      (invoice_id,amount,method,request.form.get("reference","").strip(),request.form.get("paid_on") or date.today().isoformat(),request.form.get("notes","").strip(),session["user_id"],datetime.now().isoformat(timespec="seconds")))
        flash("Payment recorded.", "success")
    except ValueError as e:
        flash(str(e), "danger")
    return redirect(url_for("invoice_detail", invoice_id=invoice_id))

@app.route("/invoices/<int:invoice_id>/pdf")
@login_required
def invoice_pdf(invoice_id):
    with db() as c:
        inv = c.execute("SELECT i.*, c.name customer,c.phone,c.address,c.gstin customer_gstin FROM invoices i JOIN customers c ON c.id=i.customer_id WHERE i.id=?", (invoice_id,)).fetchone()
        if not inv: abort(404)
        items = c.execute("SELECT * FROM invoice_items WHERE invoice_id=? ORDER BY id", (invoice_id,)).fetchall()
        payments = c.execute("SELECT * FROM payments WHERE invoice_id=? ORDER BY paid_on,id", (invoice_id,)).fetchall()
    paid = sum(p["amount_paise"] for p in payments)
    balance = max(0, inv["total_paise"] - paid)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, rightMargin=42, leftMargin=42, topMargin=38, bottomMargin=38)
    styles = getSampleStyleSheet()
    styles["Title"].fontName = "Helvetica-Bold"
    styles["Title"].fontSize = 18
    styles["Title"].leading = 22
    styles["Title"].textColor = colors.HexColor("#8b6914")
    styles["Normal"].fontName = "Helvetica"
    styles["Normal"].fontSize = 9
    styles["Normal"].leading = 13
    styles["Heading2"].fontName = "Helvetica-Bold"
    styles["Heading2"].textColor = colors.HexColor("#8b6914")
    story = [Paragraph(APP_NAME, styles["Title"]), Paragraph("Food and Catering Services", styles["Normal"])]
    if BUSINESS_ADDRESS: story.append(Paragraph(BUSINESS_ADDRESS, styles["Normal"]))
    if BUSINESS_PHONE: story.append(Paragraph("Phone: " + BUSINESS_PHONE, styles["Normal"]))
    if GSTIN: story.append(Paragraph("GSTIN: " + GSTIN, styles["Normal"]))
    story += [Spacer(1, 14), Paragraph("TAX INVOICE / BILL", styles["Heading2"]), Spacer(1, 5)]
    meta_data = [
        [Paragraph(f"<b>Invoice No.</b><br/>{inv['invoice_no']}", styles["Normal"]), Paragraph(f"<b>Invoice Date</b><br/>{inv['created_at'][:10]}", styles["Normal"])],
        [Paragraph(f"<b>Bill To</b><br/>{inv['customer']}<br/>{inv['phone'] or ''}<br/>{inv['address'] or ''}", styles["Normal"]), Paragraph(f"<b>Event</b><br/>{inv['event_name'] or '-'}<br/><b>Event Date:</b> {inv['event_date'] or '-'}", styles["Normal"])],
        [Paragraph(f"<b>Customer GSTIN:</b> {inv['customer_gstin'] or 'Not provided'}", styles["Normal"]), Paragraph(f"<b>Billing Mode:</b> {inv['billing_mode'].title()}", styles["Normal"])],
    ]
    meta = Table(meta_data, colWidths=[255, 255], hAlign="LEFT")
    meta.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#faf7ef")), ("BOX",(0,0),(-1,-1),0.6,colors.HexColor("#e5d7b5")), ("INNERGRID",(0,0),(-1,-1),0.35,colors.HexColor("#e5d7b5")), ("VALIGN",(0,0),(-1,-1),"TOP"), ("LEFTPADDING",(0,0),(-1,-1),8), ("RIGHTPADDING",(0,0),(-1,-1),8), ("TOPPADDING",(0,0),(-1,-1),7), ("BOTTOMPADDING",(0,0),(-1,-1),7)]))
    story += [meta, Spacer(1, 15)]
    data = [["Description", "Qty", "Unit", "Rate (Rs.)", "Amount (Rs.)"]]
    for it in items:
        desc = str(it["description"]).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
        data.append([Paragraph(desc, styles["Normal"]), f"{it['qty']:g}", it["unit"], f"Rs. {it['rate_paise']/100:,.2f}", f"Rs. {it['amount_paise']/100:,.2f}"])
    data += [["Subtotal", "", "", "", f"Rs. {inv['subtotal_paise']/100:,.2f}"], ["Discount", "", "", "", f"-Rs. {inv['discount_paise']/100:,.2f}"], [f"Tax ({inv['tax_rate']:g}%)", "", "", "", f"Rs. {inv['tax_paise']/100:,.2f}"], ["TOTAL", "", "", "", f"Rs. {inv['total_paise']/100:,.2f}"], ["Paid", "", "", "", f"Rs. {paid/100:,.2f}"], ["BALANCE DUE", "", "", "", f"Rs. {balance/100:,.2f}"]]
    table = Table(data, colWidths=[205, 43, 55, 92, 105], repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#c89b3c")), ("TEXTCOLOR",(0,0),(-1,0),colors.white), ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"), ("FONTSIZE",(0,0),(-1,-1),8.5), ("GRID",(0,0),(-1,-1),0.45,colors.HexColor("#dedede")), ("ALIGN",(1,1),(-1,-1),"RIGHT"), ("ALIGN",(1,0),(-1,0),"CENTER"), ("VALIGN",(0,0),(-1,-1),"MIDDLE"), ("LEFTPADDING",(0,0),(-1,-1),7), ("RIGHTPADDING",(0,0),(-1,-1),7), ("TOPPADDING",(0,0),(-1,-1),8), ("BOTTOMPADDING",(0,0),(-1,-1),8), ("BACKGROUND",(0,-3),(-1,-3),colors.HexColor("#f8f2e4")), ("FONTNAME",(0,-3),(-1,-3),"Helvetica-Bold"), ("BACKGROUND",(0,-1),(-1,-1),colors.HexColor("#fff0d2")), ("FONTNAME",(0,-1),(-1,-1),"Helvetica-Bold"), ("LINEABOVE",(0,-3),(-1,-3),1,colors.HexColor("#c89b3c"))]))
    story += [table, Spacer(1, 16)]
    if inv["notes"]:
        notes = str(inv["notes"]).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
        story += [Paragraph("<b>Notes / Terms</b>", styles["Normal"]), Paragraph(notes, styles["Normal"]), Spacer(1, 10)]
    story += [Paragraph("Thank you for choosing Arvind Pandey Food and Catering Services.", styles["Normal"]), Spacer(1, 4), Paragraph("This invoice is computer-generated.", styles["Italic"])]
    doc.build(story)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{inv['invoice_no']}.pdf", mimetype="application/pdf")

@app.route("/expenses", methods=["GET","POST"])
@login_required
def expenses():
    if request.method == "POST":
        try:
            amount = to_paise(request.form.get("amount"))
            if amount <= 0: raise ValueError("Expense amount must be greater than zero.")
            with db() as c:
                c.execute("INSERT INTO expenses(expense_date,category,description,amount_paise,payment_method,vendor,created_by,created_at) VALUES(?,?,?,?,?,?,?,?)",
                          (request.form.get("expense_date") or date.today().isoformat(),request.form.get("category","Other").strip(),request.form.get("description","").strip(),amount,request.form.get("payment_method","Cash"),request.form.get("vendor","").strip(),session["user_id"],datetime.now().isoformat(timespec="seconds")))
            flash("Expense saved.", "success")
        except ValueError as e: flash(str(e),"danger")
        return redirect(url_for("expenses"))
    with db() as c: rows = c.execute("SELECT * FROM expenses ORDER BY expense_date DESC,id DESC").fetchall()
    return render_template("expenses.html", expenses=rows)

@app.route("/inventory", methods=["GET","POST"])
@login_required
def inventory():
    if request.method == "POST":
        try:
            name = request.form.get("name","").strip()
            unit = request.form.get("unit","unit").strip()
            qty = safe_float(request.form.get("quantity","0"),0)
            reorder = safe_float(request.form.get("reorder_level","0"),0)
            cost = to_paise(request.form.get("unit_cost","0"))
            if not name: raise ValueError("Item name is required.")
            with db() as c:
                c.execute("INSERT INTO inventory(name,unit,quantity,reorder_level,unit_cost_paise,updated_at) VALUES(?,?,?,?,?,?)",
                          (name,unit,qty,reorder,cost,datetime.now().isoformat(timespec="seconds")))
            flash("Inventory item added.","success")
        except (ValueError,sqlite3.IntegrityError) as e: flash(str(e) if isinstance(e,ValueError) else "An item with this name already exists.","danger")
        return redirect(url_for("inventory"))
    with db() as c: rows=c.execute("SELECT * FROM inventory ORDER BY name").fetchall()
    return render_template("inventory.html", items=rows)

@app.route("/inventory/<int:item_id>/adjust", methods=["POST"])
@login_required
def inventory_adjust(item_id):
    try:
        delta = float(request.form.get("delta","0"))
        if delta == 0: raise ValueError("Adjustment cannot be zero.")
        reason = request.form.get("reason","Adjustment").strip() or "Adjustment"
        with db() as c:
            item=c.execute("SELECT * FROM inventory WHERE id=?",(item_id,)).fetchone()
            if not item: abort(404)
            newqty=item["quantity"]+delta
            if newqty < 0: raise ValueError("Stock cannot become negative.")
            c.execute("UPDATE inventory SET quantity=?,updated_at=? WHERE id=?",(newqty,datetime.now().isoformat(timespec="seconds"),item_id))
            c.execute("INSERT INTO stock_moves(item_id,delta,reason,note,created_by,created_at) VALUES(?,?,?,?,?,?)",(item_id,delta,reason,request.form.get("note","").strip(),session["user_id"],datetime.now().isoformat(timespec="seconds")))
        flash("Stock updated.","success")
    except ValueError as e: flash(str(e),"danger")
    return redirect(url_for("inventory"))

@app.route("/staff", methods=["GET","POST"])
@owner_required
def staff():
    if request.method == "POST":
        username=request.form.get("username","").strip()
        password=request.form.get("password","")
        if len(username)<3 or len(password)<12:
            flash("Username must be at least 3 characters and password at least 12 characters.","danger")
        else:
            try:
                with db() as c:
                    c.execute("INSERT INTO users(username,password_hash,role,created_at) VALUES(?,?,?,?)",(username,generate_password_hash(password),"staff",datetime.now().isoformat(timespec="seconds")))
                flash("Staff account created.","success")
            except sqlite3.IntegrityError: flash("Username already exists.","danger")
        return redirect(url_for("staff"))
    with db() as c: rows=c.execute("SELECT id,username,role,active,created_at FROM users ORDER BY id").fetchall()
    return render_template("staff.html", staff=rows)

@app.route("/staff/<int:user_id>/toggle", methods=["POST"])
@owner_required
def staff_toggle(user_id):
    if user_id == session["user_id"]:
        flash("You cannot deactivate your own account.","danger")
    else:
        with db() as c: c.execute("UPDATE users SET active=1-active WHERE id=?",(user_id,))
        flash("Account status updated.","success")
    return redirect(url_for("staff"))

@app.route("/reports.xlsx")
@login_required
def reports_xlsx():
    wb=Workbook()
    with db() as c:
        ws=wb.active; ws.title="Invoices"
        ws.append(["Invoice No","Date","Customer","Event","Mode","Subtotal INR","Discount INR","Tax Rate %","Tax INR","Total INR","Paid INR","Balance INR"])
        rows=c.execute("""SELECT i.*,cu.name customer,COALESCE((SELECT SUM(p.amount_paise) FROM payments p WHERE p.invoice_id=i.id),0) paid
                          FROM invoices i JOIN customers cu ON cu.id=i.customer_id ORDER BY i.id DESC""").fetchall()
        for r in rows:
            ws.append([r["invoice_no"],r["created_at"][:10],r["customer"],r["event_name"],r["billing_mode"],r["subtotal_paise"]/100,r["discount_paise"]/100,r["tax_rate"],r["tax_paise"]/100,r["total_paise"]/100,r["paid"]/100,max(0,r["total_paise"]-r["paid"])/100])
        ep=wb.create_sheet("Expenses"); ep.append(["Date","Category","Description","Vendor","Payment Method","Amount INR"])
        for r in c.execute("SELECT * FROM expenses ORDER BY expense_date DESC"):
            ep.append([r["expense_date"],r["category"],r["description"],r["vendor"],r["payment_method"],r["amount_paise"]/100])
        pp=wb.create_sheet("Payments"); pp.append(["Invoice ID","Date","Method","Reference","Amount INR"])
        for r in c.execute("SELECT * FROM payments ORDER BY paid_on DESC"):
            pp.append([r["invoice_id"],r["paid_on"],r["method"],r["reference"],r["amount_paise"]/100])
    for sheet in wb.worksheets:
        sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:
            letter=col[0].column_letter
            sheet.column_dimensions[letter].width=min(32,max(12,max(len(str(cell.value or "")) for cell in col)+2))
    buf=io.BytesIO(); wb.save(buf); buf.seek(0)
    return send_file(buf,as_attachment=True,download_name=f"catering-report-{date.today().isoformat()}.xlsx",mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/health")
def health():
    return {"status":"ok"}

@app.errorhandler(403)
def forbidden(_): return render_template("error.html",code=403,message="You do not have permission to access this page."),403
@app.errorhandler(404)
def missing(_): return render_template("error.html",code=404,message="The requested page was not found."),404

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT",5000)), debug=False)
